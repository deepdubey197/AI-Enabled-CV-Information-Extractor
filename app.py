import os
import re
import PyPDF2
from docx import Document
from openpyxl import Workbook
import requests
import streamlit as st
import zipfile

def extract_text_from_docx(file_path):
    doc = Document(file_path)
    full_text = []
    for para in doc.paragraphs:
        full_text.append(para.text)
    return '\n'.join(full_text)

def extract_text_from_pdf(file_path):
    with open(file_path, 'rb') as file:
        reader = PyPDF2.PdfFileReader(file)
        text = ''
        for page_num in range(reader.numPages):
            text += reader.getPage(page_num).extractText()
    return text

def extract_info(text):
    email_pattern = r'\b[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Z|a-z]{2,}\b'
    contact_pattern = r'\b\d{10}\b|\b\d{3}[-.\s]?\d{3}[-.\s]?\d{4}\b'

    emails = re.findall(email_pattern, text)
    contacts = re.findall(contact_pattern, text)

    return emails, contacts, text

def summarize_text(text, api_token):
    api_endpoint = "https://api-inference.huggingface.co/models/facebook/bart-large-cnn"
    headers = {
        "Authorization": f"Bearer {api_token}",
        "Content-Type": "application/json"
    }
    data = {
        "inputs": text,
        "parameters": {
            "max_length": 150,
            "min_length": 30,
            "do_sample": False
        },
        "options": {
            "use_cache": False,
            "wait_for_model": True
        }
    }
    response = requests.post(api_endpoint, headers=headers, json=data)
    if response.status_code == 200:
        summarized_text = response.json()[0]['summary_text'].strip()
        return summarized_text
    else:
        return f"Error: {response.status_code} - {response.json()['error']}"

def process_folder(folder_path, output_file):
    wb = Workbook()
    ws = wb.active
    ws.append(['Name', 'Email ID', 'Contact No.', 'Overall Text', 'Overall Summary'])

    with zipfile.ZipFile(folder_path, 'r') as zip_ref:
        zip_ref.extractall('temp_folder')

    for root, _, files in os.walk('temp_folder'):
        for file_name in files:
            file_path = os.path.join(root, file_name)

            if file_name.endswith('.docx'):
                text = extract_text_from_docx(file_path)
            elif file_name.endswith('.pdf'):
                text = extract_text_from_pdf(file_path)
            else:
                continue

            emails, contacts, full_text = extract_info(text)
            name_without_extension = os.path.splitext(file_name)[0]
            summarized_text = summarize_text(full_text, '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')

            # Append the extracted data to the worksheet
            ws.append([
                name_without_extension,
                ', '.join(emails) if emails else '',
                ', '.join(contacts) if contacts else '',
                full_text,
                summarized_text
            ])

    wb.save(output_file)
    
    # Check if 'temp_folder' exists and is empty before trying to remove it
    if os.path.exists('temp_folder') and not os.listdir('temp_folder'):
        os.rmdir('temp_folder')


def main():
    st.title('AI Enabled CV Information Extractor')
    st.markdown('This app extracts information from CVs and summarizes the overall text.')
    st.markdown('### It may take a little bit time to process the CVs because it makes use of a Large Language Model for certain tasks.')

    uploaded_folder = st.file_uploader('Upload a folder containing CVs (ZIP format)', type=['zip'])
    if uploaded_folder:
        folder_path = 'uploaded_folder.zip'
        with open(folder_path, 'wb') as f:
            f.write(uploaded_folder.read())

        output_file = 'output.xlsx'
        process_folder(folder_path, output_file)

        st.success('Processing completed! Click the button below to download the Excel file.')
        st.download_button(
            label="Download Excel file",
            data=open(output_file, 'rb'),
            file_name=output_file,
            mime='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )

if __name__ == "__main__":
    main()

